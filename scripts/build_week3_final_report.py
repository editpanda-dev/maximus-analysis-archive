from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path("/Users/woojin/Documents/막시무스")
OUT = ROOT / "output" / "reports"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "[3주차]_막시무스_방학스터디_최종보고서.docx"
CHART = ROOT / "reports" / "eda" / "dongdaemun_50min_pilot" / "figures" / "b078_time_bands.png"

NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
PALE = "EAF1F8"
LIGHT = "F3F5F7"
MID = RGBColor(92, 105, 117)
BLACK = RGBColor(20, 24, 28)
WHITE = RGBColor(255, 255, 255)
FONT_NAME = "Arial Unicode MS"


def font(run, size=8.9, bold=False, color=BLACK, italic=False):
    run.font.name = FONT_NAME
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_NAME)
    rfonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep(p, keep_next=False):
    p_pr = p._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)
    if keep_next:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, separate, end])
    font(r, 8, color=MID)


def add_header(section):
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("MAXIMUS · DAT 방학 스터디")
    font(r, 8, bold=True, color=MID)


def add_body(doc, text, size=8.9, after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    font(r, size=size)
    set_keep(p)
    return p


def add_label_paragraph(doc, label, text, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(label + "  ")
    font(r, 8.8, bold=True, color=NAVY)
    r = p.add_run(text)
    font(r, 8.8)
    set_keep(p)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    set_keep(p, keep_next=True)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    set_keep(p, keep_next=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=7.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_fill(cell, PALE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        font(r, font_size, bold=True, color=NAVY)
        if widths:
            cell.width = Mm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(text))
            font(r, font_size)
            if widths:
                cell.width = Mm(widths[i])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Mm(170)
    set_cell_fill(cell, LIGHT)
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    font(r, 9.1, bold=True, color=NAVY)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    font(r, 8.5)
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(15)
    sec.bottom_margin = Mm(14)
    sec.left_margin = Mm(19)
    sec.right_margin = Mm(19)
    sec.header_distance = Mm(7)
    sec.footer_distance = Mm(7)
    add_header(sec)
    add_page_number(sec)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal._element.rPr.rFonts.set(qn(f"w:{attr}"), FONT_NAME)
    normal.font.size = Pt(8.9)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Heading 1", 15.5, 8, 5), ("Heading 2", 11.3, 6, 3)):
        st = doc.styles[name]
        st.font.name = FONT_NAME
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            st._element.rPr.rFonts.set(qn(f"w:{attr}"), FONT_NAME)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    # Page 1 — opening and required section 1.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("[3주차] 방학 스터디 최종 보고서")
    font(r, 20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("2026년 8월 23일")
    font(r, 9.2, color=MID)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("팀 막시무스 | 장한별 · 지우진 · 최시현")
    font(r, 9.4, bold=True)

    add_h1(doc, "1. 주요 결과물 요약")
    add_callout(
        doc,
        "핵심 결론",
        "동대문구 출발 50분 생활권 추천은 ‘가까운 곳 나열’이 아니라 ① 실제 도달 가능한 후보 생성, "
        "② 유입이 소비로 이어지는 정도 추정, ③ 식사·카페·공부·여가·쇼핑·문화 목적별 적합도 계산, "
        "④ 예상보다 멀리서도 선택되는 숨은 상권 발굴의 네 단계로 설계해야 한다.",
    )
    add_h2(doc, "1.1 데이터 개요 및 검증 범위")
    add_table(
        doc,
        ["자료", "확보·검증 상태", "이번 보고서에서의 역할"],
        [
            ["B078 목적별 이동 표본", "2024.03 · 499행×47열 · 직접 EDA", "이동시간, OD 격자, 목적코드, 성·연령별 이동량"],
            ["B079 카드소비 표본", "2023.01.01 · 200행×7열 · 직접 EDA", "고객 거주지–가맹점 행정동별 업종·건수·금액"],
            ["423개 행정동 상권 EDA", "팀원 제공 결과 · 원자료 미확보", "유동·점포·업종집적·소비와 매출의 관계"],
            ["상권분석·POI 원자료", "추가 확보 필요", "목적별 시설·업종·체류 가능성 및 상권 성격"],
        ],
        widths=[38, 52, 82],
        font_size=7.5,
    )
    add_label_paragraph(doc, "직접 재현 결과", "B078 표본은 결측·중복이 없었고 이동시간 중앙값은 466초였다. B079 표본의 이용금액 합계는 약 11.36억 원, 이용건수는 36,204건, 건당 평균은 약 3.14만 원이었다. B079에는 고객 시군구 결측 17행이 존재했다.")
    add_label_paragraph(doc, "팀원 EDA 결과", "매출과 유동인구의 Spearman 상관계수는 0.555였지만 점포 수 0.899, 유사업종 점포 수 0.904, 음식지출 0.911로 더 높았다. 다만 상관은 인과가 아니며, 행정동 규모와 집계 방식이 높은 상관을 만들 수 있으므로 면적·점포규모·상권유형을 통제해야 한다.")
    add_label_paragraph(doc, "현재 한계", "B078과 B079 표본은 날짜와 공간단위가 다르고 개인 식별키도 없어서 행 단위 결합이 불가능하다. 따라서 ‘이 사람이 방문해 이 카드를 썼다’고 해석할 수 없고, 공통 기간×행정동×집단 수준에서 방문목적과 소비특성의 관계를 검증해야 한다. 또한 표본만으로 동대문구 출발 50분 후보지 전체를 복원할 수 없어 정식 원본·코드북 확보가 선행되어야 한다.")

    doc.add_page_break()
    # Page 2 — deduplicated analytical modules and EDA.
    add_h1(doc, "1. 주요 결과물 요약 (계속)")
    add_h2(doc, "1.2 중복을 제거한 최종 분석 구조")
    add_table(
        doc,
        ["통합 모듈", "통합한 기존 질문", "산출물"],
        [
            ["A. 접근성·후보군", "동대문구에서 50분 내 동네 추리기 + 출발지별 비교", "동대문구 격자·행정동 출발 50분 이하 후보 지도; 30·40·60분 민감도"],
            ["B. 소비전환·괴리", "사람이 많으면 매출도 높은가 + 통과형/목적소비형 + 숨은 상권", "기대매출 대비 잔차·달성률, 4유형 상권, 분기 안정성"],
            ["C. 목적성·추천", "공부/놀이 성격 + 멀어도 가는 곳 + 목적별 추천", "6개 목적 라벨, Hidden Score, 근거가 표시되는 순위"],
        ],
        widths=[36, 76, 60],
        font_size=7.7,
    )
    add_body(doc, "기존 10개 질문은 ‘후보 생성–괴리 측정–목적 설명 및 추천’의 세 모듈로 합쳤다. 특히 ‘유동인구 대비 매출’과 ‘예상보다 높은 숨은 상권’은 같은 괴리 분석이므로 하나의 잔차 기반 모형으로 통합한다. 선형관계에 가까운 행정동을 먼저 삭제하지 않고 전체 지역으로 기대매출을 학습한 뒤, 잔차가 작고 안정적인 지역은 일반형으로 분류하고 잔차가 크거나 비선형 패턴이 반복되는 지역만 원인 분석 대상으로 추린다.", size=8.7)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run()
    r.add_picture(str(CHART), width=Mm(169))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("그림 1. 직접 재현한 B078 전체 표본의 이동시간 구간 분포(동대문구 50분 후보지 수가 아님)")
    font(r, 7.5, color=MID)

    add_h2(doc, "1.3 EDA에서 확인한 핵심 인사이트")
    add_label_paragraph(doc, "유입 ≠ 소비", "유동인구는 잠재고객이지만 환승·통과 인구도 포함한다. 따라서 추천점수는 유입량 자체보다 실제 매출이 기대치보다 높은지, 여러 분기에 반복되는지를 함께 봐야 한다.")
    add_label_paragraph(doc, "목적별 후보는 같아도 순위는 달라야 함", "50분 필터는 모든 목적의 공통 후보군이다. 식사·카페·공부·여가·쇼핑·문화생활은 업종·POI, 시간대, 업종 다양성, B078 목적별 유입과 B079 소비 전환의 가중치가 달라 최종 순위가 달라진다.")
    add_label_paragraph(doc, "회기동 해석", "팀원 자료에 따르면 2023년 4분기 유동 약 546만 명, 점포 141개, 매출 약 87.8억 원이며 2023년 점포는 151개에서 141개로 감소했다. 대형 상권과 규모로 경쟁하기보다 대학가·식사·카페·저녁 시간대처럼 강한 이용목적을 좁혀야 한다.")
    add_label_paragraph(doc, "숨은 상권 정의", "가락1동 사례처럼 유입 규모 대비 매출이 높아 보이는 지역은 후보가 될 수 있으나, 최종 정의는 면적·점포·직장·상주인구·상권유형을 통제한 기대매출보다 실제 매출이 높고 그 초과분이 여러 분기 유지되는 지역이다.")

    doc.add_page_break()
    # Page 3 — required theoretical background.
    add_h1(doc, "2. 이론적 배경 및 심화 학습")
    add_h2(doc, "2.1 상관관계와 비선형 괴리")
    add_body(doc, "Spearman 상관계수는 각 값의 절대크기 대신 순위를 사용해 두 변수가 함께 증가·감소하는 정도를 측정한다. 선형관계가 아니어도 단조관계를 포착하고 이상치에 비교적 강하다. 그러나 ρ=0.555는 ‘유동인구가 매출의 원인’이라는 뜻이 아니며, 점포 수와 총매출의 높은 상관도 규모효과일 수 있다.")
    add_callout(doc, "괴리 측정식", "y = log(1 + 매출),  ŷ = f[log(1 + 유동), log(1 + 점포), 면적, 상권유형, 업종구성]\nGap = y − ŷ,  달성률 = 실제매출 ÷ 기대매출")
    add_body(doc, "기준선은 다중회귀로 시작하고, 유동이 늘어도 매출이 일정 수준에서 포화되는 관계는 GAM의 스플라인으로 표현한다. 그 뒤 LightGBM·CatBoost와 시간순 교차검증으로 비교한다. Gap이 지속적으로 양수면 목적소비형, 음수면 통과형 가능성이 크다. 사분면 분류는 설명용이고 최종 판정은 잔차와 분기 안정성으로 한다.")
    add_body(doc, "괴리 원인은 잔차 상·하위 행정동을 대상으로 SHAP과 부분의존도, B078의 목적·성·연령·시간대별 유입 구성, B079의 업종별 건수·금액, 환승거점·대학·시장·병원·문화시설 POI를 함께 대조해 찾는다. 이를 통해 환승·통과, 직장 점심수요, 관광·의료·교육 방문, 고단가 목적소비, 온라인 PG 매출 같은 설명가설을 만들고 분기·인접지역 비교로 확인한다.")

    add_h2(doc, "2.2 목적지 선택과 추천 단계")
    add_table(
        doc,
        ["단계", "기법", "쉽게 말하면", "활용"],
        [
            ["후보 생성", "50분 하드필터", "동대문구에서 실제로 갈 수 있는 곳만 남김", "출발지별 대중교통 접근성"],
            ["기준선", "Huff·중력모형·PPML", "규모는 클수록, 시간은 길수록 덜 선택됨", "기대 이동량과 Hidden Score"],
            ["선택 해석", "조건부/다항 로짓", "후보 간 속성이 선택확률을 어떻게 바꾸는지 계수로 설명", "이동시간·목적시설의 영향"],
            ["순위 예측", "LightGBM·CatBoost", "복잡한 비선형·상호작용을 트리 묶음으로 학습", "목적지별 이동량·순위"],
            ["설명", "SHAP", "각 변수가 예측점수를 올리거나 내린 정도", "추천 근거 카드"],
        ],
        widths=[25, 38, 61, 48],
        font_size=7.45,
    )
    add_body(doc, "다항로짓은 한 번의 선택에서 여러 목적지 중 하나를 고르는 확률을 추정한다. 후보지 j의 효용을 Uⱼ=β₁·이동시간+β₂·목적시설+β₃·다양성+…으로 놓고 exp(Uⱼ)를 모든 후보의 합으로 나눠 선택확률을 얻는다. 반복 관측과 개인차를 반영하려면 혼합로짓이 더 적합하지만, 현재 집계자료에서는 조건부로짓 또는 PPML부터 시작하는 것이 안정적이다.")

    add_h2(doc, "2.3 목적 라벨과 텍스트 개인화")
    add_body(doc, "식사는 음식점·점심/저녁 소비, 카페는 카페 밀도·체류시간, 공부는 도서관·대학·스터디시설·카페, 여가는 공원·체육·야간활동, 쇼핑은 소매점·업종다양성, 문화생활은 영화관·공연장·전시시설로 특징을 구성한다. 처음에는 전문가 규칙과 표준화 점수로 라벨을 만들고, 이후 실제 방문·선택 데이터를 학습해 가중치를 갱신한다.")
    add_body(doc, "KoRoBERTa 또는 Sentence-BERT는 유동·매출 숫자를 예측하는 주모형이 아니다. 1차 서비스가 50분 이내 행정동·상권을 추천하면, 그 안의 점포명·업종·메뉴·소개·후기 텍스트를 벡터로 바꾸어 ‘조용히 공부하기’, ‘데이트’, ‘비 오는 날 문화생활’ 같은 자연어 요구와 가까운 점포를 다시 정렬한다. 개인 로그가 쌓이면 이 벡터와 선호 이력을 MLP(다층 퍼셉트론)에 입력해 초개인화한다.")

    doc.add_page_break()
    # Page 4 — required detailed activities and contribution.
    add_h1(doc, "3. 수행 내용 상세")
    add_h2(doc, "3.1 데이터 점검·전처리·시각화")
    add_table(
        doc,
        ["수행 항목", "처리 내용", "결과·판단"],
        [
            ["스키마·품질", "CP949 인코딩, 행·열, 결측, 중복, 요약통계 확인", "B078 499×47, B079 200×7; B079 고객지역 17행 결측"],
            ["이동시간", "MOVE_TIME을 초로 가정해 분 환산; 10·20·30·40·50분 누적 비교", "중앙값 466초; 단위는 코드북으로 최종 확인"],
            ["공간 변환", "EPSG:5179 격자 중심을 위경도로 변환 후 서울 구 경계에 공간결합", "표본 중 서울 출발 184행, 동대문구 출발 5행"],
            ["B078 집계", "목적코드·성·연령·시간대별 이동량과 이동시간 분포", "목적코드 이름은 코드북 없이는 임의 해석하지 않음"],
            ["B079 집계", "업종별 이용금액·건수·건당금액, 고객거주지 분포", "PG 비중이 커 오프라인 상권평가에서는 분리 검토"],
            ["결합성 검사", "날짜·공간단위·식별키 비교", "표본 간 직접 행결합 불가; 공통 분기×행정동 집계 후 검증"],
        ],
        widths=[31, 78, 63],
        font_size=7.4,
    )
    add_h2(doc, "3.2 기준선과 분석 설계")
    add_label_paragraph(doc, "접근성 기준선", "동대문구 내부의 B078 출발 격자 또는 행정동별로 대중교통 50분 이하 목적지를 모아 합집합과 출발지별 후보를 함께 만든다. 30·40·60분에서 결과가 얼마나 바뀌는지 민감도 분석하며, 현재 표본은 동대문구 OD 전체를 포함하지 않아 정식 원본 확보 후 재실행한다.")
    add_label_paragraph(doc, "괴리 기준선", "선형관계 지역을 사전에 제거하지 않는다. log-다중회귀 → GAM → LightGBM/CatBoost 순으로 기대매출을 추정하고 시간순 검증한다. 예측잔차가 작으면 일반형, 양의 잔차가 반복되면 목적소비형, 음의 잔차가 반복되면 통과형 후보로 분류한다.")
    add_label_paragraph(doc, "추천 기준선", "거리순·인기순·Huff 점수를 비교하고, 목적별 규칙점수와 Hidden Score를 더한다. 평가는 NDCG@5·Hit Rate@5·중저인기 노출률·추천집중도와 사용자 선택률을 함께 본다.")
    add_label_paragraph(doc, "설명 기준", "SHAP은 추천점수를 올린 상위 3개 요인만 자연어로 표시한다. 인과효과가 아니라 예측근거임을 명시하고, 데이터 부족 지역에는 신뢰도 경고를 붙인다.")

    add_h2(doc, "3.3 팀원별 수행 내용")
    add_table(
        doc,
        ["팀원", "주요 수행 내용"],
        [
            ["장한별", "423개 행정동 유동–매출·점포·업종집적 상관 EDA, 가락1동·회기동 사례 및 상권 4유형 정리"],
            ["지우진", "B078/B079 표본 품질·공간·이동시간·업종 EDA, 데이터 결합 가능성 점검, 분석 파이프라인·보고서 통합"],
            ["최시현", "식사·카페·공부·여가·쇼핑·문화 목적 라벨과 POI 연결, 서비스 추천 구조 및 텍스트 재정렬 방향 설계"],
        ],
        widths=[28, 144],
        font_size=7.8,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("※ 위 분담은 현재 자료를 기준으로 정리한 보고서 편집본이며, 제출 전 팀 내부 역할 기록과 대조한다.")
    font(r, 7.3, color=MID, italic=True)

    doc.add_page_break()
    # Page 5 — required semester plan.
    add_h1(doc, "4. 차후 본 학기 계획")
    add_h2(doc, "4.1 학기 로드맵")
    add_table(
        doc,
        ["주차", "핵심 작업", "완료 기준"],
        [
            ["1–3주", "B078·B079 정식 원본/코드북, 상권분석·POI·대중교통 시간 수집; 공간·시간 단위 통일", "동대문구 OD 커버리지, 목적코드·단위 검증표, 데이터 사전"],
            ["4–6주", "R0 규칙기반 추천: 50분 후보, 목적별 시설점수, 거리·인기·Huff 기준선", "목적별 Top 5와 근거 카드; 30/40/60분 민감도"],
            ["7–9주", "유입–매출 괴리: 회귀·GAM·LightGBM/CatBoost, 시간순 검증, 잔차 안정성", "통과형·목적소비형 지도, 3개 분기 이상 재현"],
            ["10–12주", "X1 설명가능 추천: PPML/로짓, Hidden Score, SHAP, 인기편향 통제", "NDCG@5·Hidden 포함률·중저인기 노출률 비교"],
            ["13–15주", "P2 텍스트/개인화: KoRoBERTa·SBERT 재정렬, 동의 기반 행동로그와 MLP 실험", "사용자 테스트, 모델카드, 최종 데모·발표"],
        ],
        widths=[22, 94, 56],
        font_size=7.45,
    )

    add_h2(doc, "4.2 기술 고도화 단계")
    add_callout(
        doc,
        "MAX-DCR: Maximus Destination Conversion & Relevance",
        "R0 Accessible Relevance — 동대문구 출발 50분 접근성 + 목적별 규칙점수\n"
        "X1 Explainable Conversion Ranker — 이동·상권·소비전환을 학습하고 SHAP으로 설명\n"
        "P2 Perceptive Personalization — 텍스트 임베딩과 MLP로 개인 취향에 맞게 재정렬",
    )
    add_body(doc, "R0는 데이터가 적어도 작동하는 베이스 추천모델이고, X1은 실제 이동량·소비 결과로 규칙 가중치를 교정한다. P2는 개인 동의 로그가 충분할 때만 활성화한다. 개인화 데이터가 부족하거나 편향이 크면 X1의 집단 수준 추천으로 안전하게 돌아간다.", size=8.7)

    add_h2(doc, "4.3 우선순위와 의사결정 기준")
    add_table(
        doc,
        ["우선순위", "의사결정"],
        [
            ["1. 데이터", "정식 B078/B079와 코드북이 없으면 동대문구 50분 후보·목적코드 해석을 확정하지 않는다."],
            ["2. 검증", "상관계수보다 미래 분기 예측과 잔차의 지속성을 우선한다."],
            ["3. 추천", "정확도만 높이고 인기지역만 반복하는 모델은 채택하지 않는다."],
            ["4. 설명", "LLM은 데이터 기반 점수와 SHAP 근거를 문장으로 바꾸며 순위를 임의 생성하지 않는다."],
        ],
        widths=[35, 137],
        font_size=7.7,
    )
    add_body(doc, "최종적으로 본 프로젝트는 ‘사람이 많은 상권’을 추천하는 것이 아니라, 동대문구에서 50분 안에 실제로 갈 수 있고 사용 목적에 맞으며, 접근성·규모를 고려해도 방문과 소비가 기대보다 강한 행정동·상권을 먼저 제안하고, 그 안의 점포를 텍스트 적합도로 재정렬하는 것을 목표로 한다.", size=8.8, after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("자료 출처 및 재현성  ")
    font(r, 7.5, bold=True, color=NAVY)
    r = p.add_run("직접 EDA: B078_PURPOSE_250M_202403_sample.csv, B079_personal_card_inflow_dong_sample.csv. 팀원 제공 수치: 423개 행정동 EDA 요약. 전체 재현 산출물: reports/eda/dongdaemun_50min_pilot/.")
    font(r, 7.3, color=MID)

    doc.core_properties.title = "[3주차] 막시무스 방학 스터디 최종 보고서"
    doc.core_properties.subject = "동대문구 출발 50분 생활권의 목적별 상권 추천과 유입–소비 괴리 분석"
    doc.core_properties.author = "팀 막시무스 (장한별, 지우진, 최시현)"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
