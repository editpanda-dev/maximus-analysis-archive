from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/woojin/Documents/막시무스')
OUT = ROOT / 'output' / 'reports'
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / '[2주차]_막시무스_선행연구조사보고서.docx'

doc = Document()
sec = doc.sections[0]
# Named override: Korean university submission format uses A4.
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(17)
sec.bottom_margin = Mm(16)
sec.left_margin = Mm(20)
sec.right_margin = Mm(20)
sec.header_distance = Mm(8)
sec.footer_distance = Mm(8)

NAVY = RGBColor(31, 77, 120)
GRAY = RGBColor(100, 100, 100)
BLACK = RGBColor(0, 0, 0)

def font(run, size=10, bold=False, color=BLACK, name='AppleGothic'):
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{attr}'), name)
    rfonts.set(qn('w:hint'), 'eastAsia')
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color

def set_keep(p, keep_next=False, keep_lines=True):
    pPr = p._p.get_or_add_pPr()
    if keep_next:
        pPr.append(OxmlElement('w:keepNext'))
    if keep_lines:
        pPr.append(OxmlElement('w:keepLines'))

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'AppleGothic'
for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
    normal._element.rPr.rFonts.set(qn(f'w:{attr}'), 'AppleGothic')
normal.font.size = Pt(9.6)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.22
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, before, after in [('Heading 1', 14, 10, 5), ('Heading 2', 11.5, 7, 3), ('Heading 3', 10.2, 4, 2)]:
    st = styles[name]
    st.font.name = 'AppleGothic'
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        st._element.rPr.rFonts.set(qn(f'w:{attr}'), 'AppleGothic')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

def add_body(text, after=5, size=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    font(r, size or 9.6)
    set_keep(p)
    return p

def add_h1(text):
    p = doc.add_paragraph(text, style='Heading 1')
    set_keep(p, keep_next=True)
    return p

def add_h2(text):
    p = doc.add_paragraph(text, style='Heading 2')
    set_keep(p, keep_next=True)
    return p

def add_h3(text):
    p = doc.add_paragraph(text, style='Heading 3')
    set_keep(p, keep_next=True)
    return p

def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld_sep = OxmlElement('w:fldChar'); fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    font(run, 8.5, color=GRAY)

add_page_number(sec)

# Opening block: memo-style research report, intentionally compact to preserve 3-page limit.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run('[2주차] 선행 연구 조사 보고서')
font(r, 20, True, NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run('2026년 8월 16일')
font(r, 10, False, GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
r = p.add_run('팀 막시무스 | 장한별 · 지우진 · 최시현')
font(r, 10, True, BLACK)

add_h1('1. 프로젝트 조사 개요 및 방법')
add_h2('프로젝트명')
add_body('회기역 30분 생활권의 목적별 목적지 선택 분석: 수도권 생활이동·서울 카드소비·상권 특성 데이터를 활용한 조건부 상권 추천')

add_h2('프로젝트 개요')
add_body('본 프로젝트는 회기역을 출발지로 하고 대중교통 30분 이내를 공통 도달 범위로 설정한 뒤, 식사·카페·공부·여가·쇼핑·문화생활의 목적에 따라 후보 지역의 순위를 다르게 제시하는 분석을 수행한다. 30분 조건만 적용하면 모든 목적에 거의 같은 지역이 후보가 되므로, 1단계에서는 이동시간으로 갈 수 있는 지역을 만들고 2단계에서는 목적 관련 업종·시설, 업종 다양성, 시간대별 유입과 소비 특성으로 목적별 적합도를 계산한다.')
add_body('핵심 자료는 B078 수도권 목적별 이동 데이터, B079 서울 카드소비 데이터와 서울시 상권분석서비스 자료다. B078은 출발·도착 격자, 이동시간·거리, 이동목적과 성·연령별 이동량을 제공하므로 실제 이동의 중심 자료로 사용한다. B079는 개인 거래가 아니라 고객 거주 시군구와 가맹점 행정동 사이의 업종별 이용건수·금액 집계이므로 소비 결과를 확인하는 보조 자료로 활용한다. 상권분석서비스와 공공시설 자료에서는 음식점·카페·소매점, 도서관·스터디시설, 공원·체육시설, 영화관·공연장·전시시설 등의 목적별 특성을 구성한다.')

add_h2('주제 선정 동기와 조사 방법')
add_body('기존 장소 추천은 유명도나 온라인 후기 중심이어서 사용자의 출발지와 가용시간을 충분히 반영하지 못하고, 같은 30분 조건에서도 목적이 달라졌을 때 왜 추천 순위가 달라져야 하는지 설명하기 어렵다. 반면 기존 상권 연구는 매출이 높은 지역을 잘 찾지만 어느 출발지에서 어떤 경쟁 목적지 중 그곳을 선택했는지를 함께 다루지 않는 경우가 많다. 이에 접근성, 목적 적합성, 실제 이동과 소비를 하나의 구조에서 검증할 필요가 있다고 판단하였다.')
add_body('선행연구는 목적지 선택 이론, 시간 제약에 따른 후보군 구성, 다목적 상권의 흡인력, 서울 카드매출 공간분석, 인구집단별 이동 차이, 머신러닝과 이산선택모형의 결합이라는 여섯 기준으로 수집하였다. 각 연구의 데이터 단위, 후보지 구성법, 설명변수, 분석모형과 한계를 비교하고 현재 확보한 집계형 데이터에 실제 적용 가능한지를 검토하였다.', after=0)

doc.add_page_break()
add_h1('2. 주요 선행 연구 및 사례 분석')
add_h2('조사 묶음 1 | 목적지 선택의 이론·공간·개인차')
add_h3('2.1 Huff(1963) — 확률적 상권 선택')
add_body('상권 규모가 클수록 선택확률은 높고 거리저항이 커질수록 낮아진다는 확률모형을 제안했다. 여러 목적지 가운데 하나를 고르는 문제로 상권을 정의했다는 점이 핵심이다. 막시무스에서는 거리 대신 관측 이동시간을 쓰고, 매력도를 목적별 시설·업종 다양성·소비 신호로 확장한 가장 단순한 기준선으로 활용한다.', after=3, size=9.1)
add_h3('2.2 Scott & He(2012) — 시간제약 기반 선택집합')
add_body('Louisville 광역권의 쇼핑통행에 시공간 프리즘과 잠재경로구역을 적용해 실제 시간제약 안에서 도달할 수 있는 목적지만 후보로 구성하고 다항로짓을 적합했다. 모든 상권을 후보로 둔 모형보다 현실적인 선택구조를 제공하지만 표본과 지역 범위가 제한적이다. 회기역 30분 하드필터와 후보별 Long-table 구성의 직접적인 근거로 사용한다.', after=3, size=9.1)
add_h3('2.3 Arentze et al.(2005) — 다목적 쇼핑과 집적효과')
add_body('네덜란드 1,704가구의 쇼핑행동에 중첩로짓을 적용해 여러 품목을 한 번에 해결하려는 목적 조정과 상점 간 집적효과를 분석했다. 다목적성을 제외하면 대형 상권의 방문을 과소예측했다. 이에 식사·카페·쇼핑·여가 시설의 동시 존재, Shannon entropy와 HHI를 목적지 흡인력 변수로 반영한다.', after=3, size=9.1)
add_h3('2.4 양지철·이상완(2024) — 서울 카드매출 공간중심성')
add_body('서울 1,650개 상권의 카드매출에 LISA를 적용해 HH·HL·LH·LL 유형을 구분하고, 다항로지스틱 회귀와 LQ로 중심상권의 인구·토지이용 및 시간·연령별 특화를 분석했다. 출발지와 경쟁 대안은 설명하지 못한다. 막시무스에서는 LISA·LQ를 추천모형 자체가 아니라 공간맥락과 목적별 특화 변수로 사용한다.', after=3, size=9.1)
add_h3('2.5 Lenormand et al.(2015) — 사회인구학적 이동 차이')
add_body('바르셀로나와 마드리드의 4천만 건 이상 카드거래로 성별·연령·직업에 따른 이동거리, 업종, 요일·시간대 차이를 확인했다. 전체 평균만으로는 집단별 거리저항이 가려질 수 있음을 보여준다. B078에서는 성·연령과 이동시간의 상호작용을 검증하되, 개인 속성이 없는 B079 소비자료에 같은 결론을 직접 적용하지 않는다.', after=3, size=9.1)
add_h3('2.6 Alam et al.(2025) — 머신러닝과 혼합로짓 결합')
add_body('Halifax 자료에서 랜덤포레스트로 변수를 선별하고 K-prototype으로 선택집합을 만든 뒤 혼합로짓으로 선택요인의 이질성을 해석했다. 일상재와 전문상품은 이동시간 민감도가 달랐으며 소매집적 선호에도 개인차가 나타났다. 막시무스는 조건부로짓·PPML을 해석모형, LightGBM·CatBoost를 순위예측모형으로 분리하고 SHAP으로 근거를 제시한다.', after=3, size=9.1)

add_h2('조사 묶음 2 | 이동·소비 결합과 공간 추천모형')
add_h3('2.7 Klopack & Luco(2025) — 이동량과 실제 소비의 차이')
add_body('통신사 위치정보의 방문량과 결제카드 소비가 양의 상관을 보이지만 완전히 일치하지 않음을 실증했다. 유동인구가 많다는 사실만으로 지출 창출력이 높다고 단정할 수 없다는 뜻이다. B078 유입량과 B079 외부인 이용건수·금액을 분리한 뒤 유입 대비 소비전환율과 외부소비 의존도를 구성하는 근거가 된다.', after=3, size=9.1)
add_h3('2.8 Sweeney(2026) — 카드 소비의 거리감쇠와 지역 격차')
add_body('대규모 지오코딩 카드거래 흐름으로 소비 이동거리의 감쇠와 지역별 이동반경 차이를 분석했다. 사회경제적 여건에 따라 소비자가 감수하는 거리와 선택 가능한 목적지 범위가 달라질 수 있음을 보여준다. 막시무스에서는 평균 이동시간뿐 아니라 목적·시간대·출발지역 특성과 이동시간의 상호작용을 두어 동일한 30분의 의미가 집단마다 다른지 검증한다.', after=3, size=9.1)
add_h3('2.9 Cai et al.(2025) — 다중 공간단위와 2단계 추천')
add_body('격자·행정구역·세부 POI처럼 공간 해상도를 계층적으로 변환해 위치 노이즈를 줄이고 지역별 이동선호를 학습하며, 공간제약으로 후보군을 줄인 뒤 선호순위를 계산하는 파이프라인을 제시했다. 막시무스는 원자료의 격자 이동을 행정동 상권 특성과 연결할 때 교차면적·인구가중 매핑을 비교하고, 30분 후보생성과 목적별 랭킹을 분리한다.', after=3, size=9.1)
add_h3('2.10 Hounwanou et al.(2025) — 트리 기반 변수선별과 선택모형')
add_body('이동거리 외에 방문빈도, 업종구성, 교통여건 등 다차원 속성이 목적지 선택에 함께 작용함을 트리 기반 분석으로 확인하고, 조건부 추론 트리에서 고른 변수를 이산선택모형으로 연결했다. 비선형 탐색과 계수 해석을 한 모델에 억지로 맡기지 않는 접근이다. 막시무스도 트리 계열 모형으로 상호작용을 탐색한 뒤 선택·카운트모형과 결과를 대조한다.', after=3, size=9.1)

add_h2('조사 묶음 3 | Hidden Destination과 추천 편향')
add_h3('2.11 Kondo(2025) — 이동비용을 통제한 잠재 매력도')
add_body('OD 이동량, 출발·목적지 규모와 거리감쇠를 중력모형으로 분리해 “가까워서 많이 가는 곳”과 “멀어도 방문이 유지되는 곳”을 구분했다. 막시무스는 PPML로 기대 이동량을 구하고 관측-기대 잔차 또는 관측/기대 비율이 여러 기간 계속 높은 곳을 Hidden Destination으로 정의한다. 회기역 하나가 아니라 서울 전체 OD로 학습한 뒤 회기역 후보에 적용한다.', after=3, size=9.1)
add_h3('2.12 Yabe et al.(2025) — 행동 기반 장소 네트워크')
add_body('5개 미국 도시의 GPS 연속방문으로 장소 간 방향성·가중 의존망을 구축했다. 거리·POI 유형 등 물리특성은 실제 관계의 약 9~12%만 설명했고, 행동망은 충격 이후 방문변화 예측을 약 40% 개선했다. 개인 경로가 없는 B078에서는 이를 그대로 재현하지 않고 출발지×목적×시간×연령 유입벡터의 코사인 유사도로 행동 유사성 네트워크를 구성한다.', after=3, size=9.1)
add_h3('2.13 Coppolillo et al.(2024) — 인기편향과 BQS')
add_body('추천이 인기 항목을 반복 노출하는 Popularity Bias를 진단하고 정확도와 비인기 항목 품질을 함께 보는 Balanced Quality Score를 제안했다. BQS는 재정렬 알고리즘이 아니라 평가틀이다. 막시무스는 과거 유입량으로 인기군을 고정하고 Hit Rate·NDCG와 중저인기 노출률을 함께 본 뒤, 기본점수+λ×Hidden Score-γ×Popularity Penalty 재정렬의 균형점을 찾는다.', after=0, size=9.1)

doc.add_page_break()
add_h1('3. 프로젝트의 차별성')
add_h2('접근 방식의 차별화')
add_body('첫째, “회기역에서 30분”은 공통 도달 조건일 뿐 최종 추천조건이 아니다. 목적별 최소 시설을 갖춘 후보를 선별하고 업종 규모·다양성, 시간대별 유입과 소비를 적용하므로 동일한 생활권에서도 최종 순위와 추천 근거가 달라진다. 둘째, B078은 이동, B079는 집계 소비, 상권·시설 자료는 목적지 속성을 담당하게 하여 서로 다른 자료를 개인 수준으로 억지 결합하지 않는다.')
add_body('셋째, Kondo의 “이동비용 대비 초과 유입”으로 숨은 매력도를 찾고, Yabe의 관점을 변형한 행동 유사성으로 그 지역의 기능을 설명하며, Coppolillo의 평가 논리로 추천 정확도와 비인기 목적지 노출을 함께 검증한다. 즉 “많이 가는 곳 추천”이 아니라 “거리·규모를 감안해도 예상보다 선택되는 곳 발굴 → 행동 특성 확인 → 인기 편향을 통제한 노출”로 이어지는 분석 흐름이 차별점이다.')

add_h2('변수와 분석모형의 적용 계획')
add_body('종속변수는 출발조건별 목적지 이동량·점유율이며, 보조 종속변수는 카드 이용건수·금액이다. 설명변수는 이동시간·거리, 목적 관련 시설, 업종 다양성(Shannon entropy·HHI), 유동·직장·거주인구, 집객시설, LQ와 목적·연령·시간대 상호작용이다. 분석 순서는 ① 거리순·인기순·Huff 기준선 ② PPML/조건부로짓 기대 이동량과 Hidden Score ③ LightGBM·CatBoost 순위모형과 SHAP ④ 인기 패널티 재정렬 ⑤ 행동 유사성 네트워크 검증이다.')
add_body('평가는 무작위 행 분할 대신 과거기간 학습·미래기간 평가와 일부 지역 제외 평가를 사용한다. 이동량 예측은 MAE·Poisson deviance, 추천순위는 Hit Rate@k·NDCG@k, 편향은 중·저인기 노출률·추천 집중도·Hidden 포함률로 측정한다. P2 개인화는 동의받은 행동로그가 충분히 축적된 뒤 MLP로 확장하며, 데이터가 부족하면 집단 수준 X1 모델로 복귀한다.')

add_h2('기대 결과와 프로젝트의 의미')
add_body('최종 산출물은 출발지, 대중교통 최대 30분, 시간대와 목적을 입력하면 도달 가능한 후보와 목적별 순위를 제시하는 프로토타입이다. 각 결과에는 이동시간, 목적 관련 시설과 다양성, 실제 이동 및 소비 근거를 함께 표시한다. LLM은 이 근거를 자연어로 설명하는 보조 기능으로만 사용하고, 추천 순위 자체는 관측 데이터와 검증된 모형이 산출한다. 이를 통해 일반적인 장소 추천과 달리 결과의 근거, 조건별 차이와 불확실성을 확인할 수 있다.')

add_h2('참고문헌')
refs = [
    'Huff, D. L. (1963). A Probabilistic Analysis of Shopping Center Trade Areas. Land Economics, 39(1), 81-90.',
    'Scott, D. M., & He, S. Y. (2012). Modeling Constrained Destination Choice for Shopping. Journal of Transport Geography, 23, 60-71.',
    'Arentze, T. A., Oppewal, H., & Timmermans, H. J. P. (2005). A Multipurpose Shopping Trip Model. Journal of Marketing Research, 42(1), 109-115.',
    'Lenormand, M., et al. (2015). Influence of Sociodemographic Characteristics on Human Mobility. Scientific Reports, 5, 10075.',
    'Alam, M. J., Mahmud, N., & Habib, M. A. (2025). Integrating Machine Learning and Discrete Choice Modeling. Travel Behaviour and Society, 40, 100998.',
    '양지철·이상완. (2024). 카드 매출 데이터를 활용한 서울 상권 중심지의 변화 양상 및 특성 분석. 도시연구, 26, 43-75.',
    'Klopack, E., & Luco, F. (2025). JUE Insight: Measuring local consumption with payment cards and cell phone pings. Journal of Urban Economics. https://doi.org/10.1016/j.jue.2025.103798',
    'Sweeney. (2026). Spatial Variation in Mobility Across England: Evidence from 21 Million Geo-coded Credit-Card Flows. SSRN Working Paper. https://doi.org/10.2139/ssrn.5932075',
    'Cai et al. (2025). 다중 공간단위 변환 및 공간제약 기반 목적지 추천 연구. 팀 선행연구 조사자료; 최종 제출 전 원문 서지사항 확인 필요.',
    'Hounwanou et al. (2025). 트리 기반 변수선별과 이산선택 목적지 모형 연구. 팀 선행연구 조사자료; 최종 제출 전 원문 서지사항 확인 필요.',
    'Kondo, K. (2025). Measuring the attractiveness of trip destinations based on human mobility data. Scientific Reports, 15. https://doi.org/10.1038/s41598-025-29023-0',
    'Yabe, T., et al. (2025). Behaviour-based dependency networks between places shape urban economic resilience. Nature Human Behaviour, 9, 496-506. https://doi.org/10.1038/s41562-024-02072-7',
    'Coppolillo, E., et al. (2024). Balanced Quality Score: Measuring Popularity Debiasing in Recommendation. ACM Transactions on Intelligent Systems and Technology, 15(4), 1-27. https://doi.org/10.1145/3650043'
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Mm(4)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(ref)
    font(r, 7.6, color=RGBColor(55,55,55))
    set_keep(p)

doc.core_properties.title = '[2주차] 막시무스 선행 연구 조사 보고서'
doc.core_properties.subject = '회기역 30분 생활권 목적별 목적지 선택 분석'
doc.core_properties.author = '팀 막시무스 (장한별, 지우진, 최시현)'
doc.save(DOCX)
print(DOCX)
